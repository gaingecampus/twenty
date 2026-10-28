import datetime
import importlib.util
import pathlib
import unittest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

spec = importlib.util.spec_from_file_location('cash_flow', pathlib.Path(__file__).with_name('extract-cash-flow.py'))
cash = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cash)


class CashFlowExtractionTest(unittest.TestCase):
    def invoices(self):
        sheet = Workbook().active
        day = datetime.datetime(2026, 10, 1)
        for row in [
            [None, '입금 일자', '기업명', '내용', '담당', '금액'],
            [None, day, '일반', '용역', '', 100],
            [None, day, '완료', '용역', '', 200],
            [None, day, '주식회사 분할', '용역', '', 300],
            [None, '입금 일자', '기업명', '내용', '금액', '입금일', '입금액', '잔금'],
            [None, day, '별칭/분할', '용역', 300, day, 100, 200],
            [None, day + datetime.timedelta(days=30), '별칭/분할', '용역', 300, None, 0, 500],
        ]:
            sheet.append(row)
        for col in range(2, 7):
            sheet.cell(3, col).fill = PatternFill('solid', fgColor='FFFFFF00')
        return sheet

    def test_completed_and_cross_table_duplicates_excluded(self):
        result = cash.expected_incoming(self.invoices())
        self.assertEqual(result['amount'], '600')
        self.assertEqual(result['completedRows'], [3])
        self.assertEqual(result['duplicateRows'], [4])
        self.assertEqual(result['installmentClosingRows'], [7])

    def test_paid_installment_not_added_again(self):
        sheet = self.invoices()
        sheet.cell(6, 2).fill = PatternFill('solid', fgColor='FFFFFF00')
        self.assertEqual(cash.expected_incoming(sheet)['amount'], '600')

    def test_invalid_closing_balance_fails(self):
        sheet = self.invoices()
        sheet.cell(7, 8, 900)
        with self.assertRaisesRegex(ValueError, 'reconcile'):
            cash.expected_incoming(sheet)

    def test_ambiguous_overlap_fails(self):
        sheet = self.invoices()
        sheet.cell(7, 2, sheet.cell(6, 2).value)
        with self.assertRaisesRegex(ValueError, 'Ambiguous'):
            cash.expected_incoming(sheet)

    def test_latest_bank_by_time_not_row_and_zero_preserved(self):
        sheet = Workbook().active
        sheet.append([None, '2026.10.28 12:00:00', None, None, 0])
        sheet.append([None, '2026.10.27 12:00:00', None, None, 100])
        self.assertEqual(cash.latest_bank_balance(sheet, 5)['amount'], '0')

    def test_missing_latest_bank_does_not_fall_back(self):
        sheet = Workbook().active
        sheet.append([None, '2026.10.28 12:00:00'])
        sheet.append([None, '2026.10.27 12:00:00', None, None, 100])
        with self.assertRaisesRegex(ValueError, 'Missing'):
            cash.latest_bank_balance(sheet, 5)

    def test_same_time_conflicting_bank_balances_rejected(self):
        sheet = Workbook().active
        for amount in [100, 200]:
            sheet.append([None, '2026.10.28 12:00:00', None, None, amount])
        with self.assertRaisesRegex(ValueError, 'ambiguous'):
            cash.latest_bank_balance(sheet, 5)


if __name__ == '__main__':
    unittest.main()
