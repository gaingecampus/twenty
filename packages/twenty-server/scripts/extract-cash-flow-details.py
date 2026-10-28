"""Normalize private bank/receivable rows without modifying the workbook."""
import datetime
import hashlib
import importlib.util
import json
import pathlib
import re
import sys
import uuid
from decimal import Decimal

import openpyxl

spec = importlib.util.spec_from_file_location('cash_flow', pathlib.Path(__file__).with_name('extract-cash-flow.py'))
cash = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cash)


def masked(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    text = '' if value is None else str(value)
    # Preserve business labels but never copy full account/card/phone identifiers.
    return re.sub(r'(?<!\d)(?:\d[ -]?){9,19}\d(?!\d)', lambda match: '*' * (len(match[0]) - 4) + match[0][-4:], text)


def money(value, blank_zero=False):
    if value in (None, ''):
        return cash.currency(0) if blank_zero else None
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError('Non-numeric currency in source')
    amount = Decimal(str(value))
    if not amount.is_finite() or amount != amount.to_integral_value():
        raise ValueError('Invalid whole KRW amount')
    return cash.currency(amount)


def source_fields(digest, sheet, row, captured_at):
    return dict(sourceHash=digest, sourceSheet=sheet.title, sourceRow=row, capturedAt=captured_at,
                sourceUrl={'primaryLinkUrl': cash.SOURCE_URL, 'primaryLinkLabel': f'{sheet.title} {row}행', 'secondaryLinks': []})


def record_id(kind, digest, sheet, row):
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f'gainge:{kind}:{digest}:{sheet}:{row}'))


def bank_records(sheet, key, label, digest, captured_at):
    columns = {'shinhan': (4, 5, 7, 6), 'kb': (5, 4, 6, 3), 'nh': (4, 3, 5, 6)}
    incoming, outgoing, balance, description = columns[key]
    latest = cash.latest_bank_balance(sheet, balance, key == 'nh')
    records = []
    for row in range(5, sheet.max_row + 1):
        at = cash.parse_transaction_time(sheet.cell(row, 2).value, sheet.cell(row, 9).value if key == 'nh' else None)
        if at is None:
            if isinstance(sheet.cell(row, 2).value, datetime.datetime):
                raise ValueError(f'Invalid transaction time: {key} row {row}')
            continue
        bank_balance = money(sheet.cell(row, balance).value)
        if bank_balance is None:
            raise ValueError(f'Missing transaction balance: {key} row {row}')
        detail = masked(sheet.cell(row, description).value)
        records.append(dict(
            id=record_id('bankTransaction', digest, sheet.title, row),
            name=f'{label} {at:%Y-%m-%d} {detail or "거래"}',
            bankName=label, bankKey=key, transactionAt=at.isoformat() + '+09:00',
            incomingAmount=money(sheet.cell(row, incoming).value, True),
            outgoingAmount=money(sheet.cell(row, outgoing).value, True),
            balance=bank_balance, description=detail,
            sourceDetails='\n'.join(f'{masked(sheet.cell(4, col).value) or f"열 {col}"}: {masked(sheet.cell(row, col).value)}'
                                    for col in range(3, min(sheet.max_column, 11) + 1) if sheet.cell(row, col).value is not None),
            **source_fields(digest, sheet, row, captured_at),
        ))
    return records, latest


def receipt_records(sheet, digest, captured_at):
    totals = cash.expected_incoming(sheet)
    included = set(totals['ordinaryIncludedRows'])
    closing = set(totals['installmentClosingRows'])
    completed = set(totals['completedRows'])
    duplicates = set(totals['duplicateRows'])
    records, mode = [], None
    for row in range(1, sheet.max_row + 1):
        date, company, description, amount_e, amount_f, paid, balance = [sheet.cell(row, col).value for col in range(2, 9)]
        if re.sub(r'\s+', '', str(date)) == '입금일자' and company == '기업명':
            mode = 'installment' if balance == '잔금' else 'ordinary'
            continue
        if not isinstance(date, datetime.datetime) or not company:
            continue
        contribution = amount_f if row in included else balance if row in closing else 0
        status = '포함' if row in included or row in closing else '완료 제외' if row in completed else '중복 제외' if row in duplicates else '분할 중간 내역'
        records.append(dict(
            id=record_id('expectedReceipt', digest, sheet.title, row),
            name=f'{masked(company)} · {date:%Y-%m-%d}', companyName=masked(company),
            expectedDate=date.date().isoformat(), description=masked(description),
            invoiceAmount=money(amount_e if mode == 'installment' else amount_f),
            paidAmount=money(paid) if mode == 'installment' else None,
            closingBalance=money(balance) if mode == 'installment' else None,
            outstandingAmount=money(contribution), inclusionStatus=status,
            sourceDetails='\n'.join(f'열 {col}: {masked(sheet.cell(row, col).value)}' for col in range(2, 9) if sheet.cell(row, col).value is not None),
            **source_fields(digest, sheet, row, captured_at),
        ))
    actual = sum(int(record['outstandingAmount']['amountMicros']) for record in records)
    if actual != int(cash.currency(totals['amount'])['amountMicros']):
        raise ValueError('Receivable details do not reconcile')
    return records, totals


def extract(source, month):
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
        raise ValueError('Invalid month')
    workbook = openpyxl.load_workbook(source, data_only=True)
    digest = hashlib.sha256(pathlib.Path(source).read_bytes()).hexdigest()
    captured_at = datetime.datetime.fromtimestamp(pathlib.Path(source).stat().st_mtime, datetime.timezone.utc).isoformat()
    transactions, banks = [], {}
    for key, label, _ in cash.BANKS:
        sheets = [sheet for sheet in workbook if label in sheet.title and '입출금' in sheet.title]
        if len(sheets) != 1:
            raise ValueError(f'Ambiguous bank sheet: {key}')
        rows, banks[key] = bank_records(sheets[0], key, label, digest, captured_at)
        transactions.extend(rows)
    sheets = [sheet for sheet in workbook if '발행' in sheet.title and '입금' in sheet.title]
    if len(sheets) != 1:
        raise ValueError('Ambiguous invoice sheet')
    receipts, incoming = receipt_records(sheets[0], digest, captured_at)
    return dict(sourceHash=digest, capturedAt=captured_at, month=month,
                bankTransactions=transactions, expectedReceipts=receipts,
                banks=banks, expectedIncoming=cash.currency(incoming['amount']))


if __name__ == '__main__':
    result = extract(*sys.argv[1:3])
    output = pathlib.Path(sys.argv[3])
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    output.chmod(0o600)
    print(json.dumps({'bankTransactions': len(result['bankTransactions']), 'expectedReceipts': len(result['expectedReceipts'])}))
