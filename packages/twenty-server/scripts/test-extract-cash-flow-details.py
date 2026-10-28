import datetime
import importlib.util
import pathlib
import unittest
import json
from backup_identity import backup_id
from openpyxl import Workbook

def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, pathlib.Path(__file__).with_name(filename))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

details = load('details', 'extract-cash-flow-details.py')
fixtures = load('fixtures', 'test-extract-cash-flow.py')


class CashFlowDetailsTest(unittest.TestCase):
    def test_all_invoice_rows_preserved_but_only_unpaid_contributions_summed(self):
        rows, total = details.receipt_records(fixtures.CashFlowExtractionTest().invoices(), 'a' * 64, '2026-10-28T00:00:00Z')
        self.assertEqual(len(rows), 5)
        self.assertEqual(sum(int(row['outstandingAmount']['amountMicros']) for row in rows), 600_000_000)
        self.assertEqual([row['inclusionStatus'] for row in rows], ['포함', '완료 제외', '중복 제외', '분할 중간 내역', '포함'])
        self.assertEqual(rows[-1]['closingBalance']['amountMicros'], '500000000')
        self.assertEqual(total['amount'], '600')

    def test_bank_columns_and_latest_balance_reconcile(self):
        sheet = Workbook().active
        sheet.title = '신한 입출금'
        sheet.cell(4, 2, '거래일시')
        sheet.append([None, '2026.10.28 12:00:00', '123-456-789012', 50, 0, '거래', 150])
        sheet.append([None, '2026.10.27 12:00:00', None, 0, 20, '거래', 100])
        rows, latest = details.bank_records(sheet, 'shinhan', '신한', 'a' * 64, '2026-10-28T00:00:00Z')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['balance']['amountMicros'], '150000000')
        self.assertEqual(rows[0]['incomingAmount']['amountMicros'], '50000000')
        self.assertEqual(rows[1]['outgoingAmount']['amountMicros'], '20000000')
        self.assertEqual(latest['amount'], '150')
        self.assertNotIn('123-456-789012', rows[0]['sourceDetails'])
        self.assertIn('9012', rows[0]['sourceDetails'])

    def test_identity_is_stable_per_backup_and_different_between_backups(self):
        key = ('bankTransaction', 'a' * 64, '신한', 5)
        self.assertEqual(details.record_id(*key), details.record_id(*key))
        self.assertNotEqual(details.record_id(*key), details.record_id('bankTransaction', 'b' * 64, '신한', 5))

    def test_zero_preserved_missing_balance_rejected(self):
        self.assertEqual(details.money(0)['amountMicros'], '0')
        self.assertIsNone(details.money(None))
        with self.assertRaises(ValueError):
            details.money('1,000')

    def test_source_dates_are_not_mistaken_for_account_numbers(self):
        self.assertEqual(details.masked(datetime.datetime(2026, 10, 28)), '2026-10-28T00:00:00')
        self.assertNotIn('123456789012', details.masked('123456789012'))

    def test_fixed_id_survives_row_insertion_and_source_id_has_priority_over_position(self):
        sheet = Workbook().active
        sheet.title = '신한 입출금'
        sheet.cell(1, 12, 'CRM 백업 ID')
        sheet.cell(5, 12, '00000000-0000-4000-8000-000000000001')
        original = backup_id(sheet, 5)
        sheet.insert_rows(3, 2)
        self.assertEqual(backup_id(sheet, 7), original)
        sheet.cell(1, 13, '은행 거래 ID')
        sheet.cell(8, 13, 'bank-transaction-001')
        native = backup_id(sheet, 8, 'shinhan')
        sheet.insert_rows(3)
        self.assertEqual(backup_id(sheet, 9, 'shinhan'), native)

    def test_payroll_never_includes_recipient_in_name_description_or_raw_text(self):
        sheet = Workbook().active
        sheet.title = '신한 입출금'
        sheet.cell(4, 2, '거래일시')
        sheet.append([None, '2026.10.28 12:00:00', '가상직원', 0, 50, '10월 급여 가상직원', 150])
        rows, _ = details.bank_records(sheet, 'shinhan', '신한', 'a' * 64, '2026-10-28T00:00:00Z')
        self.assertNotIn('가상직원', json.dumps(rows, ensure_ascii=False))

    def test_payroll_flag_handles_source_without_salary_keyword(self):
        sheet = Workbook().active
        sheet.title = '신한 입출금'
        sheet.cell(4, 2, '거래일시')
        sheet.cell(1, 12, 'CRM 급여')
        for column, value in enumerate([None, '2026.10.28 12:00:00', '가상직원', 0, 50, '가상직원', 150], 1):
            if value is not None: sheet.cell(5, column, value)
        sheet.cell(5, 12, 'Y')
        rows, _ = details.bank_records(sheet, 'shinhan', '신한', 'a' * 64, '2026-10-28T00:00:00Z')
        self.assertNotIn('가상직원', json.dumps(rows, ensure_ascii=False))


if __name__ == '__main__':
    unittest.main()
