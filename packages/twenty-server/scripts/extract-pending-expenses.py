import sys, json, re, datetime, calendar
from decimal import Decimal
from collections import Counter
import openpyxl

source, output = sys.argv[1:]
values = openpyxl.load_workbook(source, data_only=True)
formulas = openpyxl.load_workbook(source, data_only=False)
sheet = next(s for s in values if '지출대기' in s.title)
formula_sheet = formulas[sheet.title]
records, totals, exclusions = [], [], []
year = month = None
section_status = None

def text(value):
    if value is None: return None
    if isinstance(value, datetime.datetime): return value.date().isoformat()
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip() or None

def money(value):
    if value is None: return None
    if not isinstance(value, (float, int)): raise ValueError(f'Unexpected amount {value!r}')
    return {'amountMicros': str(int(Decimal(str(value)) * 1000000)), 'currencyCode': 'KRW'}

for row_number, row in enumerate(sheet.values, 1):
    section = re.search(r'(2026)년\s*(\d+)월\)\s*(예정|확정)', str(row[1]))
    if section:
        year, month = map(int, section.group(1, 2))
        section_status = 'PLANNED' if section[3] == '예정' else 'CONFIRMED'
        continue
    if row[1] == '구분': continue
    item = text(row[4])
    if item and item.replace(' ', '') in ('총계', '소계'):
        totals.append({'row': row_number, 'month': month, 'label': item, 'values': list(row[6:10]), 'formulas': [formula_sheet.cell(row_number, c).value for c in (7, 8, 10)]})
        continue
    if not item:
        if any(v is not None for v in row[1:11]): exclusions.append(row_number)
        continue
    assert year and month
    date_value = row[8]
    due_date = None
    issue = None
    if isinstance(date_value, datetime.datetime):
        due_date = date_value.date().isoformat()
        if date_value.month != month: issue = '예정일이 귀속 월과 다름: 원본 날짜 유지'
    elif isinstance(date_value, (int, float)) and date_value == int(date_value) and 1 <= date_value <= calendar.monthrange(year, month)[1]:
        due_date = datetime.date(year, month, int(date_value)).isoformat()
    elif date_value is not None:
        issue = '예정일 확인 필요: 일자가 명확하지 않거나 해당 월에 없는 날짜'
    else:
        issue = '예정일 미입력'
    extra = '\n'.join(f'{openpyxl.utils.get_column_letter(c)}열: {text(row[c-1])}' for c in range(12, 26) if row[c-1] is not None)
    records.append({
        'name': item,
        'expenseMonth': f'{year}-{month:02}',
        'frequency': {'정기(월1회)': 'MONTHLY', '비정기': 'IRREGULAR'}.get(text(row[1])),
        'projectCategory': text(row[2]), 'accountCategory': text(row[3]), 'vendorName': text(row[5]),
        'supplyAmount': money(row[6]), 'plannedCashAmount': money(row[7]), 'paidAmount': money(row[9]),
        'scheduledDate': due_date, 'scheduledDateSource': text(date_value), 'notes': text(row[10]),
        'sourceSectionStatus': section_status, 'importReview': issue,
        'sourceRow': row_number,
        'sourceKey': f'1KfF1FAPPvD-yjAjubd1RRHew5COM5NVovgQsRsMELl8:1659969985:{row_number}',
        'sourceUrl': {'primaryLinkUrl': f'https://docs.google.com/spreadsheets/d/1KfF1FAPPvD-yjAjubd1RRHew5COM5NVovgQsRsMELl8/edit#gid=1659969985&range=B{row_number}:K{row_number}', 'primaryLinkLabel': '원본 시트 행'},
        'sourceExtraNotes': extra or None,
    })
# Preserve every remaining nonempty source row, including headings and templates.
# Summary amounts remain in the original text, not additive expense currencies.
imported_rows = {r['sourceRow'] for r in records}
year = month = None
section_status = None
for row_number, row in enumerate(sheet.values, 1):
    section = re.search(r'(2026)년\s*(\d+)월\)\s*(예정|확정)', str(row[1]))
    if section:
        year, month = map(int, section.group(1, 2))
        section_status = 'PLANNED' if section[3] == '예정' else 'CONFIRMED'
    if row_number in imported_rows or not any(v is not None for v in row):
        continue
    item = text(row[4])
    kind = ('월 제목' if section else '열 제목' if row[1] == '구분' else
            '합계·소계' if item and item.replace(' ', '') in ('총계', '소계') else
            '빈 양식' if row_number in exclusions else '추가 메모')
    raw = []
    for column, value in enumerate(row, 1):
        original = formula_sheet.cell(row_number, column).value
        if value is None and original is None: continue
        line = f'{openpyxl.utils.get_column_letter(column)}열: {text(value) or "(빈 값)"}'
        if isinstance(original, str) and original.startswith('='):
            line += f' [수식: {original}]'
        raw.append(line)
    records.append({
        'name': f'[{kind}] {item or text(row[1]) or str(row_number) + "행"}',
        'expenseMonth': f'{year}-{month:02}' if month else '2026-00',
        'frequency': None, 'projectCategory': None, 'accountCategory': None,
        'vendorName': None, 'supplyAmount': None, 'plannedCashAmount': None,
        'paidAmount': None, 'scheduledDate': None, 'scheduledDateSource': None,
        'notes': '원본 보존 행: 금액·수식은 원본 추가 메모에 보존하며 지출 합산에서 제외합니다.',
        'sourceSectionStatus': section_status, 'importReview': None,
        'sourceRecordType': kind, 'sourceRow': row_number,
        'sourceKey': f'1KfF1FAPPvD-yjAjubd1RRHew5COM5NVovgQsRsMELl8:1659969985:{row_number}',
        'sourceUrl': {'primaryLinkUrl': f'https://docs.google.com/spreadsheets/d/1KfF1FAPPvD-yjAjubd1RRHew5COM5NVovgQsRsMELl8/edit#gid=1659969985&range=B{row_number}:K{row_number}', 'primaryLinkLabel': '원본 시트 행'},
        'sourceExtraNotes': '\n'.join(raw),
    })
records.sort(key=lambda r: r['sourceRow'])
assert {r['sourceRow'] for r in records} == {
    n for n, row in enumerate(sheet.values, 1) if any(v is not None for v in row)
}
result = {'sheet': '가인지 지출대기액[2026]', 'records': records, 'sourceTotals': totals, 'preservedTemplateRows': exclusions}
with open(output, 'w') as handle: json.dump(result, handle, ensure_ascii=False, indent=2, default=str)
print(json.dumps({'count': len(records), 'byMonth': Counter(r['expenseMonth'] for r in records), 'review': Counter(r['importReview'] for r in records if r['importReview']), 'preservedTemplateRows': exclusions}, ensure_ascii=False))
