"""Read a private workbook backup; never modify the source or infer missing balances."""
import datetime
import hashlib
import json
import pathlib
import re
import sys
from decimal import Decimal
import openpyxl

SOURCE_URL = 'https://docs.google.com/spreadsheets/d/1KfF1FAPPvD-yjAjubd1RRHew5COM5NVovgQsRsMELl8/edit'
BANKS = [('shinhan', '신한', 7), ('kb', '국민', 6), ('nh', '농협', 5)]


def parse_transaction_time(value, time_value=None):
    if isinstance(value, datetime.datetime):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = datetime.datetime.strptime(value.strip(), '%Y.%m.%d %H:%M:%S')
        except ValueError:
            return None
    else:
        return None
    if time_value:
        if isinstance(time_value, datetime.time):
            parsed = datetime.datetime.combine(parsed.date(), time_value)
        else:
            try:
                parsed = datetime.datetime.combine(parsed.date(), datetime.time.fromisoformat(str(time_value).strip()))
            except ValueError:
                return None
    return parsed


def latest_bank_balance(sheet, balance_column, separate_time=False):
    transactions = []
    for row in range(1, sheet.max_row + 1):
        date = parse_transaction_time(sheet.cell(row, 2).value, sheet.cell(row, 9).value if separate_time else None)
        if date is None:
            continue
        balance = sheet.cell(row, balance_column).value
        # A newer row without a balance must not silently fall back to an older balance.
        if not isinstance(balance, (int, float)) or isinstance(balance, bool):
            transactions.append((date, row, None))
        else:
            transactions.append((date, row, Decimal(str(balance))))
    if not transactions:
        raise ValueError(f'No transactions: {sheet.title}')
    latest_time = max(item[0] for item in transactions)
    latest = [item for item in transactions if item[0] == latest_time]
    if any(item[2] is None for item in latest) or len({item[2] for item in latest}) != 1:
        raise ValueError(f'Missing or ambiguous latest balance: {sheet.title}')
    date, row, balance = latest[0]
    return {'amount': str(balance), 'at': date.isoformat() + '+09:00', 'sheet': sheet.title, 'row': row}


def currency(amount):
    return {'amountMicros': str(int(Decimal(str(amount)) * 1000000)), 'currencyCode': 'KRW'}


def company_aliases(value):
    return {re.sub(r'주식회사|\(주\)|㈜|\s+', '', part) for part in str(value).split('/')}


def expected_incoming(sheet):
    """Read both invoice tables, replacing repeated installment invoices with ONE closing balance."""
    ordinary, installments, completed = [], [], []
    mode = None
    for row in range(1, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(2, 9)]
        date, company, description, amount_e, amount_f, paid, balance = values
        if re.sub(r'\s+', '', str(date)) == '입금일자' and company == '기업명':
            mode = 'installment' if balance == '잔금' else 'ordinary'
            continue
        if not isinstance(date, datetime.datetime) or not company:
            continue
        if mode is None:
            raise ValueError(f'Invoice before header: row {row}')
        yellow = any(
            sheet.cell(row, col).fill.fgColor.type == 'rgb'
            and sheet.cell(row, col).fill.fgColor.rgb[-6:].upper() == 'FFFF00'
            for col in range(2, 7)
        )
        if yellow and mode == 'ordinary':
            completed.append(row)
            continue
        amount = amount_e if mode == 'installment' else amount_f
        if not isinstance(amount, (int, float)) or isinstance(amount, bool):
            raise ValueError(f'Missing invoice amount: row {row}')
        item = dict(row=row, date=date.date(), company=company_aliases(company),
                    amount=Decimal(str(amount)))
        if mode == 'installment':
            if not isinstance(paid, (int, float)):
                raise ValueError(f'Missing payment amount: row {row}')
            # Paid installment rows remain in the reconciliation, never in the sum.
            # Their payments have already been deducted from the closing H balance.
            item.update(paid=Decimal(str(paid)), balance=None if balance is None else Decimal(str(balance)), completed=yellow)
            installments.append(item)
        else:
            ordinary.append(item)
    if not ordinary:
        raise ValueError('No invoice rows found')
    groups = {}
    for item in installments:
        groups.setdefault(tuple(sorted(item['company'])), []).append(item)
    closing_total = Decimal(0)
    for items in groups.values():
        closing = items[-1]['balance']
        if closing is None or closing != sum(item['amount'] - item['paid'] for item in items):
            raise ValueError('Installment closing balance does not reconcile with invoices and payments')
        if items[-1]['completed'] and closing != 0:
            raise ValueError('Completed installment group has an outstanding closing balance')
        closing_total += closing
    duplicate_rows, matched_rows, ordinary_total = [], set(), Decimal(0)
    for item in ordinary:
        matches = [other for other in installments if item['company'] & other['company']
                   and item['date'] == other['date'] and item['amount'] == other['amount']]
        if len(matches) > 1 or (matches and matches[0]['row'] in matched_rows):
            raise ValueError(f'Ambiguous duplicate invoice: row {item["row"]}')
        if matches:
            duplicate_rows.append(item['row'])
            matched_rows.add(matches[0]['row'])
        else:
            ordinary_total += item['amount']
    return {
        'amount': str(ordinary_total + closing_total),
        'sheet': sheet.title,
        'completedRows': completed,
        'duplicateRows': duplicate_rows,
        'installmentClosingRows': [items[-1]['row'] for items in groups.values()],
        'ordinaryIncludedRows': [item['row'] for item in ordinary if item['row'] not in duplicate_rows],
    }


def extract(source, expenses_file, month):
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
        raise ValueError('Month must be YYYY-MM')
    workbook = openpyxl.load_workbook(source, data_only=True)
    invoice_sheets = [sheet for sheet in workbook if '발행' in sheet.title and '입금' in sheet.title]
    if len(invoice_sheets) != 1:
        raise ValueError('Ambiguous invoice sheet')
    incoming = expected_incoming(invoice_sheets[0])
    banks = {}
    for key, label, column in BANKS:
        sheets = [sheet for sheet in workbook if label in sheet.title and '입출금' in sheet.title]
        if len(sheets) != 1:
            raise ValueError(f'Ambiguous bank sheet: {label}')
        banks[key] = latest_bank_balance(sheets[0], column, key == 'nh')
    expenses = json.loads(pathlib.Path(expenses_file).read_text())
    rows = [row for row in expenses['records'] if row['expenseMonth'] == month and not row.get('sourceRecordType')]
    if not rows:
        raise ValueError(f'No expense rows for {month}')
    # H-column detail rows only; preserve blank rows in the original backup, exclude subtotals.
    expense_micros = sum(int(row['plannedCashAmount']['amountMicros']) for row in rows if row.get('plannedCashAmount') is not None)
    digest = hashlib.sha256(pathlib.Path(source).read_bytes()).hexdigest()
    captured_at = datetime.datetime.fromtimestamp(pathlib.Path(source).stat().st_mtime, datetime.timezone.utc).isoformat()
    return {
        'sourceUrl': SOURCE_URL,
        'sourceSha256': digest,
        'capturedAt': captured_at,
        'banks': banks,
        'incoming': incoming,
        'expenseDetailRows': len(rows),
        'record': {
            'name': f'{month} 자금 흐름', 'month': month,
            'capturedAt': captured_at, 'sourceHash': digest,
            'expectedIncoming': currency(incoming['amount']),
            'plannedExpense': {'amountMicros': str(expense_micros), 'currencyCode': 'KRW'},
            **{f'{key}Balance': currency(bank['amount']) for key, bank in banks.items()},
            **{f'{key}MeasuredAt': bank['at'] for key, bank in banks.items()},
            'sourceUrl': {'primaryLinkUrl': SOURCE_URL, 'primaryLinkLabel': '원본 입출금·지출대기액 시트', 'secondaryLinks': []},
        },
    }


if __name__ == '__main__':
    source, expenses, month, output = sys.argv[1:]
    result = extract(source, expenses, month)
    pathlib.Path(output).write_text(json.dumps(result, ensure_ascii=False, indent=2))
    pathlib.Path(output).chmod(0o600)
    print(json.dumps({'month': month, 'banks': len(result['banks']), 'expenseDetailRows': result['expenseDetailRows'], 'backupSha256': result['sourceSha256']}))
